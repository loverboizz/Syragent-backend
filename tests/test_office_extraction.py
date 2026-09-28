import io
import sys
sys.path.append("d:/DATN/backend")

import docx
import openpyxl
from app.services.document_processor import extract_text_from_file, split_text

# 1. Generate in-memory DOCX
doc = docx.Document()
doc.add_heading("Bảng Báo Giá Khảo Sát Drone", level=1)
doc.add_paragraph("Dịch vụ bay chụp LiDAR chính xác cao dành cho công trình xây dựng.")
table = doc.add_table(rows=2, cols=3)
table.rows[0].cells[0].text = "Tên Gói"
table.rows[0].cells[1].text = "Độ Phân Giải"
table.rows[0].cells[2].text = "Đơn Giá"
table.rows[1].cells[0].text = "Gói Tiêu Chuẩn"
table.rows[1].cells[1].text = "GSD 2cm/pixel"
table.rows[1].cells[2].text = "25.000.000 VNĐ"

docx_buf = io.BytesIO()
doc.save(docx_buf)
docx_bytes = docx_buf.getvalue()

extracted_docx = extract_text_from_file("bao_gia.docx", docx_bytes)
print("=== EXTRACTED DOCX ===")
print(extracted_docx)
assert "25.000.000 VNĐ" in extracted_docx
assert "GSD 2cm/pixel" in extracted_docx
print("[PASS] Word (.docx) extraction works perfectly!\n")

# 2. Generate in-memory XLSX
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "ThietBi"
ws.append(["Mã Dịch Vụ", "Tên Thiết Bị", "Công Suất"])
ws.append(["DV-01", "Matrice 350 RTK", "55 phút/pin"])
ws.append(["DV-02", "Zenmuse L2 LiDAR", "240.000 pts/sec"])

xlsx_buf = io.BytesIO()
wb.save(xlsx_buf)
xlsx_bytes = xlsx_buf.getvalue()

extracted_xlsx = extract_text_from_file("danh_sach_thiet_bi.xlsx", xlsx_bytes)
print("=== EXTRACTED XLSX ===")
print(extracted_xlsx)
assert "Matrice 350 RTK" in extracted_xlsx
assert "Zenmuse L2 LiDAR" in extracted_xlsx
assert "240.000 pts/sec" in extracted_xlsx
print("[PASS] Excel (.xlsx) extraction works perfectly!\n")

# 3. Test HTML extraction
html_content = b"""
<html>
<head><style>body { color: red; }</style><script>console.log('hi');</script></head>
<body>
<h1>Chinh sach bao hanh</h1>
<p>Bao hanh 12 thang doi voi thiet bi bay.</p>
</body>
</html>
"""
extracted_html = extract_text_from_file("policy.html", html_content)
print("=== EXTRACTED HTML ===")
print(extracted_html)
assert "console.log" not in extracted_html
assert "Chinh sach bao hanh" in extracted_html
print("[PASS] HTML clean extraction works perfectly!\n")

# 4. Test chunking
chunks = split_text(extracted_docx)
assert len(chunks) >= 1
print(f"[PASS] Chunking produced {len(chunks)} chunks.")

