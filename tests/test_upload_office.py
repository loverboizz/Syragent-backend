"""
Integration test for uploading Word (.docx) and Excel (.xlsx) files to an Agent's knowledge base.
"""

import asyncio
import io
import docx
import httpx
import openpyxl

BASE_URL = "http://localhost:8000"


def make_test_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Bảng Giá Dịch Vụ Bay Chụp Drone 2026", level=1)
    doc.add_paragraph("Chi tiết các gói dịch vụ trắc địa bay chụp LiDAR của GASCOLAE.")
    table = doc.add_table(rows=3, cols=3)
    table.rows[0].cells[0].text = "Gói Dịch Vụ"
    table.rows[0].cells[1].text = "Thiết Bị"
    table.rows[0].cells[2].text = "Chi Phí"
    table.rows[1].cells[0].text = "Khảo Sát Địa Hình"
    table.rows[1].cells[1].text = "DJI Matrice 350 RTK"
    table.rows[1].cells[2].text = "20.000.000 VNĐ"
    table.rows[2].cells[0].text = "Quét 3D LiDAR"
    table.rows[2].cells[1].text = "Zenmuse L2"
    table.rows[2].cells[2].text = "45.000.000 VNĐ"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_test_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BangGiaThietBi"
    ws.append(["Mã Thiết Bị", "Tên Thiết Bị", "Tần Số Quét", "Ứng Dụng"])
    ws.append(["TB-01", "Zenmuse L2", "240 kHz", "Đo vẽ mô hình 3D địa hình rừng cây dày đặc"])
    ws.append(["TB-02", "Zenmuse P1", "45 MP Full-frame", "Chụp ảnh trực thăng quang trắc đo lưới"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def test_upload_office():
    async with httpx.AsyncClient(base_url=BASE_URL, timeout=60.0) as client:
        # 1. Login
        login_res = await client.post(
            "/api/v1/auth/login",
            json={
                "email": "admin_dashboard_test@gascolae.vn",
                "password": "SecurePass123!",
            },
        )
        assert login_res.status_code == 200, f"Login failed: {login_res.text}"
        token = login_res.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}
        print("[PASS] Authenticated successfully as test admin")

        # 2. Get Services & Agent ID
        srv_res = await client.get("/api/v1/services/", headers=headers)
        assert srv_res.status_code == 200
        services = srv_res.json()
        assert len(services) > 0, "No services found for test org"
        target_service = services[0]

        # Fetch detail with attached agent
        detail_res = await client.get(f"/api/v1/services/{target_service['id']}", headers=headers)
        assert detail_res.status_code == 200
        detail_data = detail_res.json()
        if not detail_data.get("agent"):
            agent_res = await client.post(
                "/api/v1/agents/",
                headers=headers,
                json={
                    "service_id": target_service["id"],
                    "llm_provider": "google",
                    "model_name": "models/gemini-3.6-flash",
                    "ai_tier": "FREE",
                },
            )
            assert agent_res.status_code == 201, f"Failed to create agent: {agent_res.text}"
            agent_id = agent_res.json()["id"]
            print(f"[PASS] Created new Agent: {agent_id} for service '{target_service['name']}'")
        else:
            agent_id = detail_data["agent"]["id"]
            print(f"[PASS] Target Agent found: {agent_id} for service '{target_service['name']}'")

        # 3. Upload Word document
        docx_bytes = make_test_docx()
        docx_files = {
            "file": ("bang_gia_dich_vu.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        }
        res_docx = await client.post(
            f"/api/v1/agents/{agent_id}/documents/upload",
            headers=headers,
            files=docx_files,
        )
        assert res_docx.status_code == 201, f"DOCX upload failed: {res_docx.text}"
        docx_data = res_docx.json()
        print(f"[PASS] DOCX uploaded successfully! Chunks created: {docx_data['chunks_created']}")
        assert docx_data["chunks_created"] > 0

        # 4. Upload Excel document
        xlsx_bytes = make_test_xlsx()
        xlsx_files = {
            "file": ("danh_muc_thiet_bi.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        }
        res_xlsx = await client.post(
            f"/api/v1/agents/{agent_id}/documents/upload",
            headers=headers,
            files=xlsx_files,
        )
        assert res_xlsx.status_code == 201, f"XLSX upload failed: {res_xlsx.text}"
        xlsx_data = res_xlsx.json()
        print(f"[PASS] XLSX uploaded successfully! Chunks created: {xlsx_data['chunks_created']}")
        assert xlsx_data["chunks_created"] > 0

        # 5. List chunks to verify they are stored and retrievable
        list_res = await client.get(
            f"/api/v1/agents/{agent_id}/documents?limit=20",
            headers=headers,
        )
        assert list_res.status_code == 200
        chunks = list_res.json()
        sources = [c["source"] for c in chunks]
        assert "bang_gia_dich_vu.docx" in sources, "bang_gia_dich_vu.docx not found in chunks"
        assert "danh_muc_thiet_bi.xlsx" in sources, "danh_muc_thiet_bi.xlsx not found in chunks"
        print(f"[PASS] Stored chunks verified in Supabase pgvector (Found {len(chunks)} chunks)")

        print("\n[ALL TESTS PASS] Word (.docx) and Excel (.xlsx) upload and RAG vectorization work 100%!")


if __name__ == "__main__":
    asyncio.run(test_upload_office())
