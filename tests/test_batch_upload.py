"""
Integration test for Batch Multi-File Upload to an Agent's knowledge base.
Uploads multiple files (.docx, .xlsx, .txt) simultaneously in a single HTTP request.
"""

import asyncio
import io
import docx
import httpx
import openpyxl

BASE_URL = "http://localhost:8000"


def make_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Quy Trình Nghiệm Thu Dự Án 2026", level=1)
    doc.add_paragraph("Các bước nghiệm thu bàn giao bản đồ địa hình 3D và báo cáo khảo sát.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Bước"
    table.rows[0].cells[1].text = "Hạng Mục"
    table.rows[1].cells[0].text = "Bước 1"
    table.rows[1].cells[1].text = "Kiểm tra đám mây điểm Point Cloud và lưới tam giác TIN"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NhanSuDuAn"
    ws.append(["Chức Danh", "Họ Tên", "Chứng Chỉ"])
    ws.append(["Chỉ huy trưởng", "Kỹ sư Trần Văn Nam", "Chứng chỉ hành nghề đo đạc hạng I"])
    ws.append(["Phi công lái Drone", "Nguyễn Hải Đăng", "Chứng chỉ điều khiển UAV cấp bởi Cục Tác Chiến"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_txt() -> bytes:
    return "Tiêu chuẩn an toàn bay: Tuân thủ quy định bay cách xa khu vực cấm bay tối thiểu 5km. Luôn kiểm tra thời tiết và sức gió trước khi cất cánh drone.".encode("utf-8")


async def test_batch_upload():
    async with httpx.AsyncClient(base_url=BASE_URL, timeout=60.0) as client:
        # 1. Login
        login_res = await client.post(
            "/api/v1/auth/login",
            json={
                "email": "admin_dashboard_test@gascolae.vn",
                "password": "SecurePass123!",
            },
        )
        assert login_res.status_code == 200
        token = login_res.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}
        print("[PASS] Authenticated successfully as test admin")

        # 2. Get Services & Agent ID
        srv_res = await client.get("/api/v1/services/", headers=headers)
        assert srv_res.status_code == 200
        services = srv_res.json()
        assert len(services) > 0
        target_service = services[0]

        detail_res = await client.get(f"/api/v1/services/{target_service['id']}", headers=headers)
        assert detail_res.status_code == 200
        detail_data = detail_res.json()
        assert detail_data.get("agent") is not None
        agent_id = detail_data["agent"]["id"]
        print(f"[PASS] Target Agent found: {agent_id} for service '{target_service['name']}'")

        # 3. Prepare 3 files for simultaneous upload
        docx_bytes = make_docx()
        xlsx_bytes = make_xlsx()
        txt_bytes = make_txt()

        # Multi-file multipart form upload
        multi_files = [
            ("files", ("quy_trinh_nghiem_thu.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("files", ("danh_sach_nhan_su.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
            ("files", ("tieu_chuan_an_toan.txt", txt_bytes, "text/plain")),
        ]

        print(f"[INFO] Uploading 3 files simultaneously to Agent {agent_id}...")
        batch_res = await client.post(
            f"/api/v1/agents/{agent_id}/documents/upload",
            headers=headers,
            files=multi_files,
        )
        assert batch_res.status_code == 201, f"Batch upload failed: {batch_res.text}"
        batch_data = batch_res.json()

        print(f"[PASS] Batch upload response: {batch_data['message']}")
        print(f"[PASS] Total chunks created: {batch_data['total_chunks_created']}")
        assert batch_data["total_chunks_created"] >= 3
        assert len(batch_data["files_processed"]) == 3

        for f_res in batch_data["files_processed"]:
            print(f"  - {f_res['filename']}: {f_res['status']} ({f_res['chunks_created']} chunks)")
            assert f_res["status"] == "success"

        # 4. Verify all 3 files are stored in Supabase pgvector
        list_res = await client.get(
            f"/api/v1/agents/{agent_id}/documents?limit=30",
            headers=headers,
        )
        assert list_res.status_code == 200
        chunks = list_res.json()
        sources = set(c["source"] for c in chunks)

        assert "quy_trinh_nghiem_thu.docx" in sources
        assert "danh_sach_nhan_su.xlsx" in sources
        assert "tieu_chuan_an_toan.txt" in sources
        print(f"[PASS] All 3 file sources verified in Supabase vector store!")

        print("\n[ALL TESTS PASS] Batch multi-file upload is 100% operational!")


if __name__ == "__main__":
    asyncio.run(test_batch_upload())

